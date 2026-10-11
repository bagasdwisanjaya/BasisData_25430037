import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="999999"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def markdown_to_docx(md_path, docx_path):
    base_dir = os.path.dirname(os.path.abspath(md_path))
    doc = Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return
        
        # Parse table lines
        parsed_rows = []
        for tline in table_lines:
            tline = tline.strip()
            if not tline.startswith('|') or not tline.endswith('|'):
                continue
            cells = [c.strip() for c in tline.split('|')[1:-1]]
            # check if separator row
            if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                continue
            parsed_rows.append(cells)
        
        if parsed_rows:
            num_cols = max(len(r) for r in parsed_rows)
            table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(table)

            for row_idx, row_data in enumerate(parsed_rows):
                row = table.rows[row_idx]
                is_header = (row_idx == 0)
                for col_idx in range(num_cols):
                    cell_text = row_data[col_idx] if col_idx < len(row_data) else ""
                    cell = row.cells[col_idx]
                    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                    if is_header:
                        set_cell_background(cell, "2B4C7E")
                    elif row_idx % 2 == 1:
                        set_cell_background(cell, "F7F9FC")
                    else:
                        set_cell_background(cell, "FFFFFF")
                    
                    # Clean markdown bold inside cells
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    
                    parts = re.split(r'(\*\*.*?\*\*)', cell_text)
                    for pt in parts:
                        if pt.startswith('**') and pt.endswith('**'):
                            run = p.add_run(pt[2:-2])
                            run.bold = True
                        else:
                            run = p.add_run(pt)
                        if is_header:
                            run.font.color.rgb = RGBColor(255, 255, 255)
                            run.bold = True
                        else:
                            run.font.color.rgb = RGBColor(30, 30, 30)
                        run.font.size = Pt(9.5)
                        run.font.name = 'Calibri'
        
        # Add space after table
        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        table_lines = []
        in_table = False

    def flush_code():
        nonlocal in_code_block, code_lines
        if not code_lines:
            in_code_block = False
            return
        code_text = "".join(code_lines)
        
        # Render code box as a single-cell shaded table
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        set_cell_background(cell, "23272E")
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(code_text.rstrip())
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(220, 225, 230)
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        code_lines = []
        in_code_block = False

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle code fence
        if stripped.startswith('```'):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle table
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            flush_table()

        # Empty lines
        if not stripped:
            i += 1
            continue

        # Horizontal rules
        if stripped in ('---', '***', '___'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # Headings
        if stripped.startswith('# '):
            h = doc.add_heading(stripped[2:].strip(), level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            for run in h.runs:
                run.font.name = 'Calibri'
                run.font.color.rgb = RGBColor(24, 43, 73)
            i += 1
            continue
        elif stripped.startswith('## '):
            h = doc.add_heading(stripped[3:].strip(), level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            for run in h.runs:
                run.font.name = 'Calibri'
                run.font.color.rgb = RGBColor(35, 75, 125)
            i += 1
            continue
        elif stripped.startswith('### '):
            h = doc.add_heading(stripped[4:].strip(), level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            for run in h.runs:
                run.font.name = 'Calibri'
                run.font.color.rgb = RGBColor(50, 90, 140)
            i += 1
            continue
        elif stripped.startswith('#### '):
            h = doc.add_heading(stripped[5:].strip(), level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            for run in h.runs:
                run.font.name = 'Calibri'
            i += 1
            continue

        # Images: ![caption](path)
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            caption, img_rel_path = img_match.groups()
            img_rel_path = img_rel_path.strip()
            img_full_path = os.path.join(base_dir, img_rel_path) if img_rel_path else ""
            
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)

            if img_full_path and os.path.exists(img_full_path):
                try:
                    p.add_run().add_picture(img_full_path, width=Inches(5.5))
                except Exception as e:
                    run = p.add_run(f"[Gambar: {caption} ({img_rel_path})]")
                    run.font.italic = True
            else:
                run = p.add_run(f"📷 [Area Tangkapan Layar: {caption}]")
                run.font.bold = True
                run.font.color.rgb = RGBColor(120, 120, 120)
            
            if caption:
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_after = Pt(8)
                run_cap = p_cap.add_run(f"Gambar: {caption}")
                run_cap.font.italic = True
                run_cap.font.size = Pt(8.5)
                run_cap.font.color.rgb = RGBColor(100, 100, 100)
            i += 1
            continue

        # Blockquote: > text
        if stripped.startswith('> '):
            quote_text = stripped[2:].strip()
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            set_cell_background(cell, "F0F4F8")
            set_cell_margins(cell, top=60, bottom=60, left=120, right=100)
            
            # Left blue bar
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="0" w:color="2B4C7E"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
            tcPr.append(tcBorders)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(quote_text)
            run.font.italic = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(50, 70, 90)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # Numbered list: 1. text
        num_list_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_list_match:
            num_str, item_text = num_list_match.groups()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            
            # Parse bold/italic in list item
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', item_text)
            for pt in parts:
                if pt.startswith('**') and pt.endswith('**'):
                    run = p.add_run(pt[2:-2])
                    run.bold = True
                elif pt.startswith('*') and pt.endswith('*'):
                    run = p.add_run(pt[1:-1])
                    run.font.italic = True
                elif pt.startswith('`') and pt.endswith('`'):
                    run = p.add_run(pt[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9)
                else:
                    p.add_run(pt)
            i += 1
            continue

        # Bullet list: - text or * text
        bullet_list_match = re.match(r'^[-*]\s+(.*)$', stripped)
        if bullet_list_match:
            item_text = bullet_list_match.group(1)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', item_text)
            for pt in parts:
                if pt.startswith('**') and pt.endswith('**'):
                    run = p.add_run(pt[2:-2])
                    run.bold = True
                elif pt.startswith('*') and pt.endswith('*'):
                    run = p.add_run(pt[1:-1])
                    run.font.italic = True
                elif pt.startswith('`') and pt.endswith('`'):
                    run = p.add_run(pt[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9)
                else:
                    p.add_run(pt)
            i += 1
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', stripped)
        for pt in parts:
            if pt.startswith('**') and pt.endswith('**'):
                run = p.add_run(pt[2:-2])
                run.bold = True
            elif pt.startswith('*') and pt.endswith('*'):
                run = p.add_run(pt[1:-1])
                run.font.italic = True
            elif pt.startswith('`') and pt.endswith('`'):
                run = p.add_run(pt[1:-1])
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
            else:
                p.add_run(pt)
        i += 1

    if in_code_block:
        flush_code()
    if in_table:
        flush_table()

    doc.save(docx_path)
    print(f"Successfully generated: {docx_path}")

if __name__ == "__main__":
    import sys
    if len(sys.argv) >= 3:
        markdown_to_docx(sys.argv[1], sys.argv[2])
    else:
        print("Usage: python convert_to_docx.py input.md output.docx")
